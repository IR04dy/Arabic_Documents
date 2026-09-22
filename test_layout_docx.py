# -*- coding: utf-8 -*-
"""Formatted Word export (layout_docx.py) against the sample documents.

The layouts below are shaped like the ones /extract returns (Surya's blocks:
label, box as page fractions, the text lines each contributed); their boxes
were read off the pages at 200 DPI. They mix the shapes Surya produces for a
form — one block per row, label and value as separate lines of one block, a
pipe table, blocks side by side — so every path through the exporter runs.

Sizes depend on whether Pillow can shape Arabic (libraqm) and an Arial file is
found, so they are asserted as ranges, never exactly.
"""
import io
import unittest
from types import SimpleNamespace

import docx
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph

import layout_docx as L


def _box(w, h, label, x0, y0, x1, y1, *lines):
    return {"label": label, "bbox": [x0 / w, y0 / h, x1 / w, y1 / h], "lines": list(lines)}


# 01_expired_wakalah.pdf — a two-page Najiz power of attorney (1653x2339 px)
_W = lambda *a: _box(1653, 2339, *a)
_HDR = ("المملكة العربية السعودية", "وزارة العدل — كتابة العدل",
        "بوابة ناجز الإلكترونية · وكالة إلكترونية مصدّقة")
WAKALAH = [
    {"blocks": [
        _W("PageHeader", 600, 112, 1050, 232, *_HDR),
        _W("SectionHeader", 716, 296, 934, 350, "وكالة شرعية"),
        _W("SectionHeader", 1318, 395, 1512, 446, "بيانات الوكالة"),
        _W("Text", 712, 476, 1518, 520, "رقم الوكالة ٤٣٨٢١٩٠٠٥١١٢"),
        _W("Text", 732, 540, 1518, 580, "تاريخ الوكالة ١٤٤٥/٠١/١٥هـ"),
        _W("Text", 834, 603, 1518, 643, "حالة الوكالة سارية"),
        _W("Text", 732, 665, 1518, 705, "تاريخ انتهاء الوكالة ١٤٤٧/٠٣/٢٩هـ"),
        _W("Text", 494, 726, 1518, 770, "الجهة المصدرة كتابة العدل الأولى بمدينة الرياض"),
        _W("SectionHeader", 1330, 822, 1512, 872, "بيانات الموكل"),
        _W("Text", 520, 902, 1518, 1133, "الاسم", "محمد بن عبدالله بن سالم السالم", "رقم الهوية",
           "١٠٢٣٤٥٦٧٨٩", "الجنسية", "سعودي", "تاريخ انتهاء الهوية", "١٤٥٢/٠٨/١١هـ"),
        _W("SectionHeader", 1322, 1186, 1512, 1236, "بيانات الوكيل"),
        _W("Table", 500, 1268, 1518, 1438, "الاسم | خالد بن سعد بن مبارك القحطاني",
           "رقم الهوية | ١٠٤٤٥٥٦٦٧٧", "الجنسية | سعودي"),
        _W("SectionHeader", 1200, 1487, 1512, 1537, "موضوع الوكالة ونوعها"),
        _W("Text", 714, 1568, 1518, 1612, "موضوع الوكالة بيع وإفراغ عقار"),
        _W("Text", 830, 1631, 1518, 1671, "نوع الوكالة خاصة"),
    ]},
    {"blocks": [
        _W("PageHeader", 600, 112, 1050, 232, *_HDR),
        _W("SectionHeader", 1058, 300, 1512, 350, "نص الوكالة والصلاحيات الممنوحة"),
        _W("Text", 1222, 383, 1538, 427, "المادة الأولى: نص الوكالة"),
        _W("Text", 110, 439, 1540, 540, "أقر الموكل المذكور أعلاه بأنه وكّل وكيله المذكور وكالة شرعية "
           "معتبرة في بيع وإفراغ عقار، وفي مراجعة كافة الجهات الحكومية والأهلية ذات العلاقة، والتوقيع "
           "نيابةً عنه فيما ذُكر، وقبض الثمن وتسليمه، واستلام الصكوك والوثائق اللازمة لذلك."),
        _W("Text", 890, 574, 1518, 614, "حق توكيل الغير لا"),
        _W("Text", 1100, 770, 1518, 815, "توقيع الموكل: .............................."),
        _W("Text", 378, 770, 844, 815, "ختم كاتب العدل: .............................."),
        _W("PageFooter", 382, 904, 1270, 946, "وثيقة صادرة إلكترونيًا ولا تحتاج إلى ختم يدوي · "
           "للتحقق من صحة الوثيقة يُرجى زيارة بوابة ناجز"),
    ]},
]

# 01_shisha_cafe_license_request.pdf — a vector form: QR code, ruled tables (1655x2340)
_S = lambda *a: _box(1655, 2340, *a)
SHISHA = [{"blocks": [
    _S("SectionHeader", 1040, 198, 1462, 258, "طلب إصدار رخصة بلدية"),
    _S("Text", 1040, 266, 1462, 318, "مقهى يقدّم منتجات التبغ (المعسل)"),
    _S("Picture", 196, 196, 324, 328),
    _S("Text", 510, 383, 1450, 462, "نموذج اختباري",
       "أُعدّ هذا المستند لاختبار الاسترجاع في قاعدة البيانات. بيانات مقدّم الطلب والمنشأة افتراضية بالكامل"),
    _S("SectionHeader", 1176, 509, 1462, 558, "أولاً: بيانات مقدّم الطلب"),
    _S("Table", 196, 569, 1460, 1016, "البند | البيان", "الاسم | (احمد احمد بن عدوان)",
       "رقم الهوية الوطنية | ١٢٣٤٥٦٧٨٩", "الصفة | مالك المنشأة", "رقم الجوال | ٠٥٠٥٠٥٠٥٠٥٠٥",
       "البريد الإلكتروني | applicant@example.test", "السن عند تقديم الطلب | ٤٥"),
    _S("SectionHeader", 1216, 1050, 1462, 1096, "ثانياً: بيانات المنشأة"),
    _S("Table", 196, 1109, 1460, 1619, "البند | البيان", "الاسم التجاري | مقهى الرياض",
       "رقم السجل التجاري | ١٠١٠١٠١٠١٠١٠", "النشاط المطلوب | تشغيل مقهى + تقديم منتجات التبغ (المعسل)",
       "المدينة / الحي | الرياض — حي الشفا", "المساحة | ٢٤٠ م٢ (مائتان وأربعون متراً مربعاً)",
       "نوع المبنى ورقمه | تجاري — رقم المبنى ٢٣", "صفة الانتفاع | عقد إيجار مدته ثلاث سنوات"),
    _S("SectionHeader", 1046, 1652, 1462, 1698, "ثالثاً: النشاط المطلوب الترخيص له"),
    _S("Text", 570, 1711, 1462, 1757, "يُطلب الترخيص بتقديم المعسل للمرتادين داخل النطاق العمراني، بوصف المنشأة «مقهى»."),
    _S("SectionHeader", 1395, 1795, 1462, 1840, "إقرار"),
    _S("Text", 198, 1852, 1462, 1898, "يقرّ مقدّم الطلب بصحة البيانات، وبالالتزام بكامل اشتراطات لائحة "
       "تقديم منتجات التبغ، وبأن أي مخالفة تعرّضه للعقوبات النظامية."),
    _S("Table", 196, 1914, 1460, 2048, "الاسم | التوقيع | تاريخ انتهاء الطلب",
       "احمد احمد بن عدوان | احمد بن عدوان | ٢١/٩/١٤٤٣"),
]}]

# Sak.pdf — a scanned heirs deed: white text on filled bands and cells (1700x2200)
_K = lambda *a: _box(1700, 2200, *a)
_HEIRS = [
    "الابن | ناصر نافع سبيل الحربي | السعودية | ١٠٠٢٩٥٢٨٥٩ | ١٤٠٥/٠٣/٢٥ هـ | راشد",
    "الزوجة | بدريه علي سليم الحربي | السعودية | ١٠٠١١٦٧٦٥٧ | ١٣٨٤/٠٧/٠١ هـ | راشد",
    "الزوجة | عميشاء تريحيب ساجر الحربي | السعودية | ١٠١١٣٨٧٢٦١ | ١٣٨٧/٠٧/٠١ هـ | راشد",
    "الابنة | شيخه نافع سبيل الحربي | السعودية | ١٠٠١١٦٧٦٦٥ | ١٣٩٧/٠٧/٠١ هـ | راشد",
    "الابن | بدر نافع سبيل الحربي | السعودية | ١٠٠١١٦٧٦٨١ | ١٤٠٣/٠٤/٠٢ هـ | راشد",
    "الابن | منصور نافع سبيل الحربي | السعودية | ١٠٥٠٢٠٤٣٩٣ | ١٤٠٦/١١/٢٠ هـ | راشد",
    "الابنة | مزنه نافع سبيل الحربي | السعودية | ١٠٠١١٦٧٦٧٣ | ١٤٠١/١٢/٢٥ هـ | راشد",
    "الابن | عبدالمحسن نافع سبيل الحربي | السعودية | ١٠٥٠٢٠٤٤٠١ | ١٤٠٧/٠٥/١٨ هـ | راشد",
    "الابن | فهد نافع سبيل الحربي | السعودية | ١٠٩٨٠٧٩٧١٦ | ١٤١٨/٠٨/٠٦ هـ | راشد",
    "الابن | فايز نافع سبيل الحربي | السعودية | ١٠٦٠٩٧٣٤٥٨ | ١٤٠٨/١٢/١٢ هـ | راشد",
]
SAK = [{"blocks": [
    _K("Picture", 1130, 25, 1505, 168),
    _K("Text", 625, 88, 1075, 160, "وثيقة ورثة متوفى"),
    _K("Text", 262, 88, 625, 160, "كتابة العدل"),
    _K("Picture", 52, 68, 258, 252),
    _K("Text", 290, 170, 630, 325, "رقم الوثيقة : ٤٣١٥٣٧٢٣٤", "تاريخها : ١٤٤٣/٠٤/١٩ هـ",
       "الموافق : ٢٠٢١/١١/٢٤ م"),
    _K("SectionHeader", 715, 410, 910, 460, "تم توثيق ورثة"),
    _K("Table", 128, 533, 1497, 857, "اسم المتوفى | نافع سبيل عوده الحربي", "جنس المتوفى | ذكر",
       "نوع الهوية | هوية وطنية", "رقم الهوية | ١٠٠١١٦٧٦٤٠", "تاريخ الوفاة | ١٤٤٣/٠٤/٠٩ هـ",
       "رقم شهادة الوفاة | ١٨٦٥-٠٠٠٠٢١٣٤"),
    _K("Text", 1255, 980, 1500, 1025, "وانحصر ورثته في"),
    _K("Table", 133, 1052, 1502, 1932,
       "صلة القرابة | الاسم | الجنسية | الهوية الوطنية | تاريخ الميلاد | حالة الوارث", *_HEIRS),
    _K("Text", 55, 1980, 245, 2015, "1 من 2"),
    _K("Text", 135, 2040, 1335, 2080, "صدرت هذه الوثيقة من النظام الإلكتروني لخدمات التوثيق، ويمكن "
       "التحقق من صحة الوثيقة عبر الخدمات الإلكترونية لوزارة العدل"),
    _K("Picture", 1332, 2008, 1442, 2122),
]}]


def page_texts(layouts):
    return ["\n".join(t for b in p["blocks"] for t in b["lines"]) for p in layouts]


def export(pdf, layouts=None, pages=None):
    with open(pdf, "rb") as f:
        data = f.read()
    pages = pages or page_texts(layouts)
    clean = L.clean_layouts(layouts, len(pages)) if layouts is not None else None
    return docx.Document(io.BytesIO(L.build_layout_docx(data, pages, "x", layouts=clean)))


def body(d):
    """[(kind, object)] in document order: ("p", Paragraph) or ("t", Table)."""
    out = []
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            out.append(("p", Paragraph(el, d)))
        elif tag == "tbl":
            out.append(("t", DocxTable(el, d)))
    return out


def para(d, text):
    return next(p for k, p in body(d) if k == "p" and p.text == text)


def ppr(p, name):
    pPr = p._p.pPr
    return None if pPr is None else pPr.find(qn("w:" + name))


def cell_fill(cell):
    tcPr = cell._tc.tcPr
    shd = None if tcPr is None else tcPr.find(qn("w:shd"))
    return None if shd is None else shd.get(qn("w:fill"))


def size(p):
    return p.runs[0].font.size.pt


def colour(p):
    c = p.runs[0].font.color
    return str(c.rgb) if c is not None and c.type is not None else None


class LayoutValidation(unittest.TestCase):
    def test_a_layout_must_be_a_list(self):
        with self.assertRaises(ValueError):
            L.clean_layouts({"blocks": []}, 1)

    def test_a_malformed_page_is_dropped_whole(self):
        good = {"blocks": [{"label": "Text", "bbox": [0.1, 0.1, 0.5, 0.2], "lines": ["نص"]}]}
        bad_box = {"blocks": [{"label": "Text", "bbox": [0.1, "x", 0.5, 0.2], "lines": ["نص"]}]}
        empty_box = {"blocks": [{"label": "Text", "bbox": [0.5, 0.1, 0.5, 0.2], "lines": ["نص"]}]}
        self.assertEqual(L.clean_layouts([good, bad_box, empty_box, None], 4),
                         [good, None, None, None])

    def test_boxes_are_clamped_and_extra_pages_ignored(self):
        page = {"blocks": [{"label": "Text", "bbox": [-1, 0.1, 2, 0.2], "lines": []}]}
        self.assertEqual(L.clean_layouts([page, page], 1),
                         [{"blocks": [{"label": "Text", "bbox": [0.0, 0.1, 1.0, 0.2], "lines": []}]}])


class TextMatching(unittest.TestCase):
    BLOCKS = [{"label": "Text", "bbox": [0, 0, 1, 1], "lines": ["الاسم محمد", "رقم الهوية ١٢٣"]},
              {"label": "Text", "bbox": [0, 0, 1, 1], "lines": ["نص الوكالة"]}]

    def test_identical_lines_pair_exactly(self):
        self.assertEqual(L._assign(self.BLOCKS, ["الاسم محمد", "رقم الهوية ١٢٣", "نص الوكالة"]),
                         [(0, 0), (0, 1), (1, 0)])

    def test_a_proofread_line_keeps_its_place(self):
        self.assertEqual(L._assign(self.BLOCKS, ["الإسم محمد", "رقم الهوية ١٢٣", "نص الوكاله"]),
                         [(0, 0), (0, 1), (1, 0)])

    def test_an_inserted_line_hangs_on_the_line_before(self):
        self.assertEqual(L._assign(self.BLOCKS, ["الاسم محمد", "سطر جديد", "رقم الهوية ١٢٣", "نص الوكالة"]),
                         [(0, 0), (0, 0), (0, 1), (1, 0)])


class ExtractLayout(unittest.TestCase):
    """extract._page_text keeps what Surya found alongside the text."""

    def run_page(self, blocks):
        import extract
        from PIL import Image
        rec = lambda imgs, full_page: [SimpleNamespace(blocks=blocks)]
        return extract._page_text(rec, Image.new("RGB", (1000, 2000), "white"))

    def block(self, label, order, html, bbox=(100, 200, 900, 260), **kw):
        return SimpleNamespace(label=label, reading_order=order, html=html, bbox=list(bbox),
                               skipped=kw.get("skipped", False), error=False)

    def test_blocks_carry_their_lines_and_pictures_their_place(self):
        text, layout = self.run_page([
            self.block("SectionHeader", 0, "<h2>وكالة شرعية</h2>"),
            self.block("Picture", 1, "", (10, 10, 110, 110), skipped=True),
            self.block("Table", 2, "<table><tr><td>الاسم</td><td>محمد</td></tr></table>"),
        ])
        self.assertEqual(text, "وكالة شرعية\nالاسم | محمد")
        self.assertEqual([b["label"] for b in layout["blocks"]], ["SectionHeader", "Picture", "Table"])
        self.assertEqual(layout["blocks"][1]["lines"], [])
        self.assertEqual(layout["blocks"][0]["bbox"], [0.1, 0.1, 0.9, 0.13])
        self.assertEqual("\n".join(t for b in layout["blocks"] for t in b["lines"]), text)

    def test_a_text_block_without_a_box_voids_the_layout(self):
        text, layout = self.run_page([self.block("Text", 0, "<p>نص</p>", bbox=(5, 5, 5, 5))])
        self.assertEqual((text, layout), ("نص", None))


class TextOnlyPath(unittest.TestCase):
    """A page with no layout: the fixes to the old heuristics."""

    @classmethod
    def setUpClass(cls):
        L.USE_PALETTE = False          # no Surya layout model in the test environment
        cls.d = export("01_expired_wakalah.pdf", pages=page_texts(WAKALAH))

    def test_the_title_is_centred_without_a_sampled_colour(self):
        p = para(self.d, "وكالة شرعية")
        self.assertEqual(ppr(p, "jc").get(qn("w:val")), "center")
        self.assertEqual(size(p), L.TITLE_PT)

    def test_a_section_heading_over_a_table_stays_a_heading(self):
        p = para(self.d, "بيانات الوكالة")
        self.assertEqual(size(p), L.HEADING_PT)
        self.assertEqual(ppr(p, "outlineLvl").get(qn("w:val")), "1")

    def test_an_unknown_label_and_its_date_join_the_table(self):
        rows = [[c.text for c in r.cells] for k, t in body(self.d) if k == "t" for r in t.rows]
        self.assertIn(["تاريخ انتهاء الهوية", "١٤٥٢/٠٨/١١هـ"], rows)

    def test_colons_are_kept(self):
        texts = [p.text for k, p in body(self.d) if k == "p"]
        cells = [c.text for k, t in body(self.d) if k == "t" for r in t.rows for c in r.cells]
        self.assertIn("المادة الأولى: نص الوكالة", texts)
        self.assertIn("توقيع الموكل:", cells)

    def test_row_rules(self):
        self.assertIsNone(L._cells("المادة الأولى: نص الوكالة"))            # not a declared label
        self.assertIsNone(L._is_field("توقيع الموكل: …… ختم كاتب العدل: ……"))  # two fields, one line
        self.assertEqual(L._trailing_value("تاريخ انتهاء الهوية ١٤٥٢/٠٨/١١هـ"),
                         ["تاريخ انتهاء الهوية", "١٤٥٢/٠٨/١١هـ"])
        self.assertTrue(L._mostly_digits("١٤٥٢/٠٨/١١هـ"))                     # Arabic-Indic digits


class WakalahFromLayout(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.d = export("01_expired_wakalah.pdf", WAKALAH)
        cls.items = body(cls.d)

    def test_page_setup_is_the_originals(self):
        sec = self.d.sections[0]
        self.assertAlmostEqual(sec.page_width.pt, 595, delta=1)
        self.assertAlmostEqual(sec.page_height.pt, 842, delta=1)

    def test_letterhead_and_title_are_centred(self):
        for text in _HDR + ("وكالة شرعية",):
            self.assertEqual(ppr(para(self.d, text), "jc").get(qn("w:val")), "center", text)
        self.assertEqual(ppr(para(self.d, "وكالة شرعية"), "outlineLvl").get(qn("w:val")), "0")
        self.assertEqual(colour(para(self.d, _HDR[2])), "575757")          # the grey third line

    def test_section_headings_are_shaded_with_their_accent_bar(self):
        for text in ("بيانات الوكالة", "بيانات الموكل", "بيانات الوكيل", "موضوع الوكالة ونوعها"):
            p = para(self.d, text)
            self.assertEqual(ppr(p, "shd").get(qn("w:fill")), "EEF1F5", text)
            self.assertIsNotNone(ppr(p, "pBdr").find(qn("w:right")), text)
            self.assertTrue(p.runs[0].bold, text)
            self.assertEqual(ppr(p, "outlineLvl").get(qn("w:val")), "1", text)
        sizes = {size(para(self.d, t)) for t in ("بيانات الوكالة", "بيانات الموكل", "بيانات الوكيل")}
        self.assertEqual(len(sizes), 1)                                     # one level, one size

    def test_every_form_shape_becomes_a_label_value_table(self):
        tables = [t for k, t in self.items if k == "t"]
        rows = [[c.text for c in r.cells] for t in tables for r in t.rows]
        for row in (["رقم الوكالة", "٤٣٨٢١٩٠٠٥١١٢"],                      # one block per row
                    ["الاسم", "محمد بن عبدالله بن سالم السالم"],             # label, value lines
                    ["تاريخ انتهاء الهوية", "١٤٥٢/٠٨/١١هـ"],                # a label no template knows
                    ["الاسم", "خالد بن سعد بن مبارك القحطاني"],              # a pipe table
                    ["حق توكيل الغير", "لا"]):
            self.assertIn(row, rows)

    def test_labels_are_quiet_and_values_strong(self):
        t = next(t for k, t in self.items if k == "t")
        label, value = (c.paragraphs[0] for c in t.rows[0].cells)
        self.assertFalse(label.runs[0].bold)
        self.assertEqual(colour(label), "3D4753")
        self.assertTrue(value.runs[0].bold)
        self.assertEqual({size(c.paragraphs[0]) for r in t.rows for c in r.cells[1:]},
                         {size(value)})                                     # one column, one size

    def test_page_two(self):
        clause = para(self.d, "المادة الأولى: نص الوكالة")
        self.assertTrue(clause.runs[0].bold)
        prose = next(p for k, p in self.items if k == "p" and p.text.startswith("أقر الموكل"))
        self.assertEqual(ppr(prose, "jc").get(qn("w:val")), "both")
        signatures = next(t for k, t in self.items if k == "t" and "توقيع" in t.rows[0].cells[0].text)
        self.assertEqual([c.text for c in signatures.rows[0].cells],
                         ["توقيع الموكل: ..............................",
                          "ختم كاتب العدل: .............................."])
        footer = para(self.d, WAKALAH[1]["blocks"][-1]["lines"][0])
        self.assertEqual(ppr(footer, "jc").get(qn("w:val")), "center")
        breaks = [p for k, p in self.items if k == "p" and ppr(p, "pageBreakBefore") is not None]
        self.assertEqual([p.text for p in breaks], [_HDR[0]])              # page 2 opens on its letterhead

    def test_rules_are_drawn(self):
        rules = [p for k, p in self.items if k == "p" and not p.text
                 and ppr(p, "pBdr") is not None and ppr(p, "pBdr").find(qn("w:bottom")) is not None]
        self.assertEqual(len(rules), 3)          # under each letterhead, above the footer

    def test_sizes_are_sane(self):
        for k, p in self.items:
            if k == "p" and p.runs and p.text:
                self.assertTrue(8 <= size(p) <= 18, (p.text, size(p)))


class ShishaFromLayout(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.d = export("01_shisha_cafe_license_request.pdf", SHISHA)
        cls.items = body(cls.d)

    def test_the_qr_code_sits_beside_the_title(self):
        first = next(t for k, t in self.items if k == "t")
        right, left = first.rows[0].cells
        self.assertEqual([p.text for p in right.paragraphs],
                         ["طلب إصدار رخصة بلدية", "مقهى يقدّم منتجات التبغ (المعسل)"])
        self.assertTrue(left.paragraphs[0]._p.xpath(".//w:drawing"))

    def test_headings_keep_their_green(self):
        self.assertEqual(colour(para(self.d, "أولاً: بيانات مقدّم الطلب")), "1F6B4A")

    def test_the_note_box_is_one_shaded_band(self):
        a, b = para(self.d, SHISHA[0]["blocks"][3]["lines"][0]), para(self.d, SHISHA[0]["blocks"][3]["lines"][1])
        self.assertEqual(ppr(a, "shd").get(qn("w:fill")), ppr(b, "shd").get(qn("w:fill")))
        self.assertEqual(ppr(a, "pBdr").find(qn("w:top")).get(qn("w:space")),
                         ppr(b, "pBdr").find(qn("w:top")).get(qn("w:space")))

    def test_ruled_tables_keep_their_rules_header_and_columns(self):
        t = next(t for k, t in self.items if k == "t" and t.rows[0].cells[0].text == "البند")
        borders = t._tbl.tblPr.find(qn("w:tblBorders"))
        self.assertEqual(borders.find(qn("w:insideH")).get(qn("w:val")), "single")
        self.assertIsNotNone(cell_fill(t.rows[0].cells[0]))                  # shaded header row
        self.assertIsNone(cell_fill(t.rows[1].cells[0]))
        self.assertTrue(t.rows[0].cells[0].paragraphs[0].runs[0].bold)
        grid = [int(c.get(qn("w:w"))) for c in t._tbl.find(qn("w:tblGrid"))]
        self.assertAlmostEqual(grid[0] / sum(grid), 0.33, delta=0.03)     # from the rule at x≈1040

    def test_regular_text_is_not_bold(self):
        rows = {r.cells[0].text: r.cells[1].paragraphs[0] for k, t in self.items if k == "t"
                for r in t.rows if len(r.cells) == 2}
        self.assertFalse(rows["البريد الإلكتروني"].runs[0].bold)
        self.assertFalse(rows["الاسم"].runs[0].bold)


class SakFromLayout(unittest.TestCase):
    """A scan, with white text on filled bands and cells."""

    @classmethod
    def setUpClass(cls):
        cls.d = export("Sak.pdf", SAK)
        cls.tables = [t for k, t in body(cls.d) if k == "t"]

    def test_filled_bands_become_shading_with_light_text(self):
        band = next(p for t in self.tables for r in t.rows for c in r.cells for p in c.paragraphs
                    if p.text == "كتابة العدل")
        fill = ppr(band, "shd").get(qn("w:fill"))
        self.assertTrue(int(fill[2:4], 16) > int(fill[0:2], 16))           # green
        self.assertTrue(all(int(colour(band)[i:i + 2], 16) > 200 for i in (0, 2, 4)))  # white text
        orange = next(p for t in self.tables for r in t.rows for c in r.cells for p in c.paragraphs
                      if p.text == "وثيقة ورثة متوفى")
        self.assertIsNotNone(ppr(orange, "shd"))

    def test_filled_label_cells(self):
        t = next(t for t in self.tables if t.rows[0].cells[0].text == "اسم المتوفى")
        self.assertEqual(len(t.rows), 6)
        for r in t.rows:
            self.assertIsNotNone(cell_fill(r.cells[0]))
            self.assertIsNone(cell_fill(r.cells[1]))

    def test_the_heirs_table(self):
        t = next(t for t in self.tables if t.rows[0].cells[0].text == "صلة القرابة")
        self.assertEqual((len(t.rows), len(t.columns)), (11, 6))
        self.assertIsNotNone(cell_fill(t.rows[0].cells[0]))                 # dark header row
        self.assertEqual(t.rows[3].cells[1].text, "عميشاء تريحيب ساجر الحربي")
        grid = [int(c.get(qn("w:w"))) for c in t._tbl.find(qn("w:tblGrid"))]
        self.assertEqual(len(set(grid)), len(grid))                        # columns from the rules

    def test_it_stays_one_page(self):
        breaks = [p for k, p in body(self.d) if k == "p" and ppr(p, "pageBreakBefore") is not None]
        self.assertEqual(breaks, [])


if __name__ == "__main__":
    unittest.main()
